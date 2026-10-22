"""Word documents and a workbook for tests of text blocks whose values come from Word.

``wordDocument`` builds a minimal .docx in memory. ``externalValuesWorkbook`` is the 1.3.0 sample
with ``EXTERNAL_CONCEPT``, a text block it reports, named in a ``template_external_values`` range,
since no shipped template uses one.
"""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl.workbook.defined_name import DefinedName

from mireport.xlsx_template_reader.util import loadExcelFromPathOrFileLike

SAMPLE_1_3_0 = (
    Path(__file__).parents[2]
    / "digital-templates"
    / "VSME-Digital-Template-Sample-1.3.0.xlsx"
)
EXTERNAL_CONCEPT = "vsme:DescriptionOfATargetRelatedToAPolicy"

_CONTENT_TYPES = (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)


def wordDocument(
    text: str, *, paragraphStyle: str | None = None, breakType: str | None = None
) -> bytes:
    """A .docx of one paragraph. A ``paragraphStyle`` is not defined in the document, so mammoth
    warns about it; a ``breakType`` it doesn't support, mammoth warns about and drops."""
    style = (
        f'<w:pPr><w:pStyle w:val="{paragraphStyle}"/></w:pPr>' if paragraphStyle else ""
    )
    lineBreak = f'<w:br w:type="{breakType}"/>' if breakType else ""
    document = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p>{style}<w:r><w:t>{text}</w:t>{lineBreak}</w:r></w:p></w:body></w:document>"
    )
    buffer = BytesIO()
    with ZipFile(buffer, "w") as docx:
        docx.writestr("[Content_Types].xml", _CONTENT_TYPES)
        docx.writestr("_rels/.rels", _RELS)
        docx.writestr("word/document.xml", document)
    return buffer.getvalue()


def externalValuesWorkbook() -> bytes:
    workbook = loadExcelFromPathOrFileLike(SAMPLE_1_3_0)
    sheet = workbook.worksheets[0]
    cell = sheet.cell(row=1, column=250)  # far out, otherwise unused
    cell.value = EXTERNAL_CONCEPT.partition(":")[2]
    workbook.defined_names.add(
        DefinedName(
            "template_external_values", attr_text=f"'{sheet.title}'!{cell.coordinate}"
        )
    )
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
