"""Last year's report as a filled-in Digital Template, read by PriorReport for its comparatives."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from mireport.comparatives import PriorReport
from mireport.conversionresults import ConversionResultsBuilder, Stage
from mireport.data.disclosures import VSME_DEFAULTS
from mireport.exceptions import EarlyAbortException
from mireport.filesupport import FilelikeAndFileName
from mireport.report.comparatives import SkipReason
from mireport.xlsx_template_reader.processor import XlsxProcessor

SAMPLES = Path("digital-templates")


NOT_SAVED_IN_EXCEL = (
    r"^last-year\.xlsx has been migrated but not yet saved in "
    r"Excel\. Open it in Excel, save it, and give it again\.$"
)


def _migratedNotSaved() -> Workbook:
    wb = Workbook()
    wb.defined_names["template_migration_status"] = DefinedName(
        "template_migration_status", attr_text="Sheet!$A$1"
    )
    return wb


def test_a_migrated_workbook_not_yet_saved_in_excel_is_refused() -> None:
    """Its formula cells, the migration status among them, have no values until Excel saves it."""
    unsaved = BytesIO()
    _migratedNotSaved().save(unsaved)
    with pytest.raises(EarlyAbortException, match=NOT_SAVED_IN_EXCEL):
        PriorReport.fromWorkbook(
            FilelikeAndFileName(unsaved.getvalue(), "last-year.xlsx"),
            ConversionResultsBuilder(),
        )


def test_the_workbook_is_had_from_the_loader_given() -> None:
    """A workbook already loaded (to check it for migration, say) isn't parsed again."""
    loaded: list[str] = []

    def loadWorkbook(source: FilelikeAndFileName) -> Workbook:
        loaded.append(source.filename)
        return _migratedNotSaved()

    with pytest.raises(EarlyAbortException, match=NOT_SAVED_IN_EXCEL):
        PriorReport.fromFile(
            FilelikeAndFileName(b"never parsed", "last-year.xlsx"),
            ConversionResultsBuilder(),
            arelle=None,  # type: ignore[arg-type]
            loadWorkbook=loadWorkbook,
        )
    assert loaded == ["last-year.xlsx"]


def test_a_loader_that_cant_read_it_says_it_isnt_a_workbook() -> None:
    def loadWorkbook(source: FilelikeAndFileName) -> Workbook:
        raise BadZipFile("File is not a zip file")

    with pytest.raises(
        EarlyAbortException,
        match=r"^last-year\.xlsx is not an Excel workbook: File is not a zip file$",
    ):
        PriorReport.fromWorkbook(
            FilelikeAndFileName(b"never parsed", "last-year.xlsx"),
            ConversionResultsBuilder(),
            loadWorkbook=loadWorkbook,
        )


@pytest.mark.slow
def test_last_years_workbook_on_an_older_template_gives_the_comparatives() -> None:
    """The 1.2.0 sample, on the taxonomy before ours, is for the year before the 1.3.0 one's."""
    current = XlsxProcessor.from_file(
        SAMPLES / "VSME-Digital-Template-Sample-1.3.0.xlsx",
        ConversionResultsBuilder(consoleOutput=False),
        VSME_DEFAULTS,
    ).createReport()
    results = ConversionResultsBuilder(consoleOutput=False)
    path = SAMPLES / "VSME-Digital-Template-Sample-1.2.0.xlsx"
    prior = PriorReport.fromWorkbook(
        FilelikeAndFileName(path.read_bytes(), path.name), results
    )
    assert prior.report.taxonomy is not current.taxonomy

    merged = prior.mergeInto(current, results)
    assert merged.taken > 50
    assert set(merged.skipped) == {
        SkipReason.NOT_COMPARABLE,
        SkipReason.OTHER_PERIOD,
        SkipReason.NOTHING_THIS_YEAR,
    }
    assert str(current.priorPeriod) == "2025-01-01T00:00:00/2026-01-01T00:00:00"
    # Only reading last year's workbook warns (it is an older template), not the merge.
    reading, merging = results.sections
    assert reading.stage is merging.stage is Stage.PRIOR
    assert reading.hasWarnings()
    assert not merging.hasWarnings()
    current.requireAllFactsRendered = True
    assert current.getInlineReport().fileContent
