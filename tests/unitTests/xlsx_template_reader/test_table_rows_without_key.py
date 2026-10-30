"""A table row with values but an empty dimension cell gives one warning, on
that cell, listing what was not reported; rows holding only FALSE (unticked
checkboxes) stay quiet."""

from datetime import date

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from mireport.conversionresults import ConversionResultsBuilder, MessageType, Severity
from mireport.data.disclosures import VSME_DEFAULTS
from mireport.report import InlineReport
from mireport.taxonomy import getTaxonomy
from mireport.xlsx_template_reader._binder import WorkbookBinder
from mireport.xlsx_template_reader._fact_creator import FactCreator
from mireport.xlsx_template_reader._reader import WorkbookReader

VS_ENTRY_POINT = "https://xbrl.efrag.org/taxonomy/vs/2026-10-16/vs-all.xsd"
SITES = {
    "SitesInBiodiversitySensitiveAreasTable": "A",
    "IdentifierOfSiteTypedAxis_SitesInBiodiversitySensitiveAreasTable": "A",
    "SiteLocatedInABiodiversitySensitiveArea": "B",
    "SiteLocatedNearABiodiversitySensitiveArea": "C",
}
WASTE = {
    "WasteGeneratedTable": "A",
    "TypeOfWasteAxis": "A",
    "WasteGeneratedMass": "B",
}


def _warnings(
    columns: dict[str, str], rows: list[list[object]]
) -> list[tuple[str, str | None]]:
    """Convert a table (first name in *columns*, the rest its columns) holding
    these rows from row 2, and return the conversion warnings with their Excel refs."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    last = 1 + len(rows)
    table, *names = columns
    lastCol = max(columns[n] for n in names)
    wb.defined_names[table] = DefinedName(table, attr_text=f"S!$A$2:${lastCol}${last}")
    for name in names:
        col = columns[name]
        wb.defined_names[name] = DefinedName(
            name, attr_text=f"S!${col}$2:${col}${last}"
        )

    results = ConversionResultsBuilder(consoleOutput=False)
    reader = WorkbookReader(wb, results)
    taxonomy = getTaxonomy(VS_ENTRY_POINT)
    bindings = WorkbookBinder(reader, taxonomy, results).bind()
    assert bindings.tables
    report = InlineReport(taxonomy, None)
    report.addDurationPeriod("cur", date(2025, 1, 1), date(2025, 12, 31))
    report.setDefaultPeriodName("cur")
    FactCreator(bindings, reader, report, results, VSME_DEFAULTS).create_all_facts()
    return [
        (str(m.messageText), m.excelReference)
        for m in results.messages
        if m.severity is Severity.WARNING and m.messageType is MessageType.Conversion
    ]


def test_row_without_site_id_gives_one_warning_on_the_id_cell():
    assert _warnings(SITES, [["S1", True, False], [None, True, False]]) == [
        (
            (
                'Row 3 of "Sites in biodiversity sensitive areas" has no "Identifier of site", '
                'so its values were not reported: "Site located in a biodiversity sensitive area" = TRUE, '
                '"Site located near a biodiversity sensitive area" = FALSE.'
            ),
            "'S'!$A$3",
        )
    ]


def test_row_holding_only_false_is_quiet():
    assert _warnings(SITES, [[None, False, False], [None, False, False]]) == []


def test_zero_is_data():
    assert _warnings(WASTE, [[None, 0]]) == [
        (
            (
                'Row 2 of "Waste generated" has no "Type of waste" selected, '
                'so its values were not reported: "Waste generated (mass)" = 0.'
            ),
            "'S'!$A$2",
        )
    ]
