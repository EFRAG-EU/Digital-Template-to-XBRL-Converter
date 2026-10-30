from collections.abc import Generator

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from mireport.xlsx_template_reader._ranges import (
    CellRangeMetadata,
    _CellRangeDimensions,
    getEffectiveCellRangeDimensions,
)


@pytest.fixture
def sample_worksheet() -> Generator[Worksheet, None, None]:
    wb: Workbook = Workbook()
    ws: Worksheet = wb.active

    # Fill sample data
    ws["A1"] = "Header1"
    ws["B1"] = "Header2"
    ws["A2"] = "Data1"
    ws["B2"] = None
    ws["A3"] = None
    ws["B3"] = "Data2"
    ws["C1"] = None  # Intentionally blank

    yield ws


def test_filled_and_empty_cells(sample_worksheet: Worksheet) -> None:
    cr: CellRange = CellRange("A1:C3")
    dims: _CellRangeDimensions = getEffectiveCellRangeDimensions(sample_worksheet, cr)

    assert dims.countPopulated == 4
    assert dims.populated_width == 2  # Column C is completely empty
    assert dims.populated_height == 3


def test_completely_empty_range(sample_worksheet: Worksheet) -> None:
    cr: CellRange = CellRange("D1:E5")
    dims: _CellRangeDimensions = getEffectiveCellRangeDimensions(sample_worksheet, cr)

    assert dims.countPopulated == 0
    assert dims.populated_width == 1  # min width enforced
    assert dims.populated_height == 1  # min height enforced


def test_partially_filled_rows(sample_worksheet: Worksheet) -> None:
    sample_worksheet["C2"] = "ExtraData"
    cr: CellRange = CellRange("A1:C3")
    dims: _CellRangeDimensions = getEffectiveCellRangeDimensions(sample_worksheet, cr)

    assert dims.countPopulated == 5
    assert dims.populated_width == 3  # All columns have data now
    assert dims.populated_height == 3


def test_single_cell_range(sample_worksheet: Worksheet) -> None:
    cr: CellRange = CellRange("A1")
    dims: _CellRangeDimensions = getEffectiveCellRangeDimensions(sample_worksheet, cr)

    assert dims.countPopulated == 1
    assert dims.populated_width == 1
    assert dims.populated_height == 1


def test_several_areas_count_no_cells_between_them(
    sample_worksheet: Worksheet,
) -> None:
    sample_worksheet["A5"] = "Gap"
    dims = getEffectiveCellRangeDimensions(
        sample_worksheet, CellRange("A1:B1"), CellRange("A9:B9")
    )

    assert len(dims.cellsAccessed) == 4
    assert dims.countPopulated == 2


def _crm(ws: Worksheet, *areas: str) -> CellRangeMetadata:
    first, *more = (CellRange(a) for a in areas)
    return CellRangeMetadata(
        DefinedName("X", attr_text=""),
        ws,
        first,
        populated_width=1,
        populated_height=1,
        populated_min_col=1,
        populated_min_row=1,
        moreAreas=tuple(more),
    )


class TestSeveralAreas:
    def test_contains_only_what_is_inside_an_area(
        self, sample_worksheet: Worksheet
    ) -> None:
        table = _crm(sample_worksheet, "A1:C3", "A9:C9")
        assert table.contains(_crm(sample_worksheet, "B2"))
        assert table.contains(_crm(sample_worksheet, "C9"))
        assert not table.contains(_crm(sample_worksheet, "B5"))
        assert not table.contains(_crm(sample_worksheet, "A3:A9"))

    def test_overlaps_any_area(self, sample_worksheet: Worksheet) -> None:
        table = _crm(sample_worksheet, "A1:C3", "A9:C9")
        assert table.overlaps(_crm(sample_worksheet, "C8:C9"))
        assert not table.overlaps(_crm(sample_worksheet, "A5:C5"))

    def test_excel_ref_lists_every_area(self, sample_worksheet: Worksheet) -> None:
        table = _crm(sample_worksheet, "A1:C3", "A9:C9")
        assert table.excelRef() == "'Sheet'!$A$1:$C$3,'Sheet'!$A$9:$C$9"
