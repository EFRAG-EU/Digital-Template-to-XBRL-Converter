"""A Concept_Member cell is reported the same whether it is a simple fact or a
primary item in a table: with its member's dimension, and in its member's
named period (a range named for the member, such as BaselineYearMember,
holding a year)."""

from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from mireport.conversionresults import ConversionResultsBuilder
from mireport.data.disclosures import VSME_DEFAULTS
from mireport.report import InlineReport
from mireport.taxonomy import getTaxonomy
from mireport.xlsx_template_reader._binder import WorkbookBinder
from mireport.xlsx_template_reader._fact_creator import FactCreator
from mireport.xlsx_template_reader._reader import WorkbookReader

VS_ENTRY_POINT = "https://xbrl.efrag.org/taxonomy/vs/2026-10-16/vs-all.xsd"
SCOPE1 = "GrossScope1GreenhouseGasEmissions"
REUSE = "PercentageOfWasteDivertedToPreparationForReuseMass"
W01 = "W-01-WastesResultingFromExplorationMiningQuarryingAndPhysicalAndChemicalTreatmentOfMineralsMember"


def _facts(table: str | None, cells: dict[str, tuple[str, object]]) -> dict[str, str]:
    """Convert these named cells (name -> (cell, value)), inside a table range
    over all of them when table is given, and return each fact's concept and
    members with its period's name."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for name, (cell, value) in cells.items():
        ws[cell] = value
        wb.defined_names[name] = DefinedName(name, attr_text=f"S!${cell[0]}${cell[1:]}")
    if table is not None:
        wb.defined_names[table] = DefinedName(table, attr_text="S!$A$1:$C$3")

    results = ConversionResultsBuilder(consoleOutput=False)
    reader = WorkbookReader(wb, results)
    taxonomy = getTaxonomy(VS_ENTRY_POINT)
    bindings = WorkbookBinder(reader, taxonomy, results).bind()
    assert bool(bindings.tables) is (table is not None)
    report = InlineReport(taxonomy, None)
    report.addDurationPeriod("cur", date(2025, 1, 1), date(2025, 12, 31))
    report.setDefaultPeriodName("cur")
    FactCreator(bindings, reader, report, results, VSME_DEFAULTS).create_all_facts()
    return {
        " ".join(
            [f.concept.qname.localName]
            + [m.qname.localName for m in f.explicit_dimensions.values()]
        ): f.period.reportPeriod.name
        for f in report.facts
    }


@pytest.mark.parametrize("inTable", [False, True], ids=["simple", "table"])
def test_member_cell_takes_its_members_period(inTable: bool):
    facts = _facts(
        "EstimatedAbsoluteGreenhouseGasEmissionsTable" if inTable else None,
        {
            "BaselineYearMember": ("B2", 2020),
            f"{SCOPE1}_BaselineYearMember": ("B3", 150),
            SCOPE1: ("C3", 100),
        },
    )
    assert facts == {
        f"{SCOPE1} BaselineYearMember": "BaselineYearMember",
        SCOPE1: "cur",
    }


@pytest.mark.parametrize("inTable", [False, True], ids=["simple", "table"])
def test_member_of_a_dimension_without_default_is_set(inTable: bool):
    facts = _facts(
        "WasteGeneratedTable" if inTable else None, {f"{REUSE}_{W01}": ("B2", 0.5)}
    )
    assert facts == {f"{REUSE} {W01}": "cur"}
