"""Tests for the WorkbookReader public API: resolveRange, value()/CellValue,
and the WorkbookBinder that turns a workbook into WorkbookBindings."""

from datetime import date, datetime
from pathlib import Path
from typing import ClassVar

import pytest
from openpyxl.workbook.defined_name import DefinedName

from mireport.conversionresults import ConversionResultsBuilder
from mireport.data.disclosures import VSME_DEFAULTS
from mireport.taxonomy import getTaxonomy
from mireport.xlsx_template_reader._binder import WorkbookBinder
from mireport.xlsx_template_reader._constants import ALL_ERROR_VALUES
from mireport.xlsx_template_reader._ranges import (
    CellRangeMetadata,
    XbrlConceptCellRangeMetadata,
)
from mireport.xlsx_template_reader._reader import CellValue, WorkbookReader
from mireport.xlsx_template_reader.util import loadExcelFromPathOrFileLike

SAMPLE = (
    Path(__file__).parent.parent.parent
    / "data"
    / "VSME-Digital-Template-Sample-1.2.0.xlsx"
)


def _results() -> ConversionResultsBuilder:
    return ConversionResultsBuilder(consoleOutput=False)


@pytest.fixture(scope="module")
def reader():
    wb = loadExcelFromPathOrFileLike(SAMPLE)
    yield WorkbookReader(wb, _results())
    wb.close()


@pytest.fixture()
def fresh_reader():
    """Function-scoped reader for tests that mutate used-name tracking."""
    wb = loadExcelFromPathOrFileLike(SAMPLE)
    yield WorkbookReader(wb, _results())
    wb.close()


@pytest.fixture(scope="module")
def taxonomy(reader):
    entry_point = reader.value(VSME_DEFAULTS["entryPoint"]).as_str()
    return getTaxonomy(entry_point)


class TestResolveRange:
    def test_known_name_resolves(self, reader):
        crm = reader.resolveRange("template_reporting_template_version")
        assert isinstance(crm, CellRangeMetadata)
        assert crm.definedName.name == "template_reporting_template_version"

    def test_missing_name_returns_none(self, reader):
        assert reader.resolveRange("this_does_not_exist_xyz") is None

    def test_accepts_defined_name(self, reader):
        dn = reader.getDefinedName("template_reporting_template_version")
        assert dn is not None
        crm = reader.resolveRange(dn)
        assert isinstance(crm, CellRangeMetadata)

    def test_accepts_cell_range_metadata_passthrough(self, reader):
        crm = reader.resolveRange("template_reporting_template_version")
        assert reader.resolveRange(crm) is crm

    def test_marks_name_as_used(self, fresh_reader):
        resolved_dn = None
        for dn in sorted(fresh_reader.unused_defined_names, key=lambda d: d.name):
            if fresh_reader.resolveRange(dn) is not None:
                resolved_dn = dn
                break
        assert resolved_dn is not None, "no resolvable defined name found in sample"
        assert resolved_dn not in fresh_reader.unused_defined_names

    def test_peek_does_not_mark_used(self, fresh_reader):
        peeked_dn = None
        for dn in sorted(fresh_reader.unused_defined_names, key=lambda d: d.name):
            if fresh_reader.peekRange(dn) is not None:
                peeked_dn = dn
                break
        assert peeked_dn is not None, "no resolvable defined name found in sample"
        assert peeked_dn in fresh_reader.unused_defined_names


class TestCellValue:
    def test_none_is_blank(self):
        assert CellValue(None).isBlank

    def test_placeholder_dash_is_blank(self):
        assert CellValue("-").isBlank

    @pytest.mark.parametrize("error_value", sorted(ALL_ERROR_VALUES))
    def test_error_values_are_not_blank(self, error_value):
        """isBlank and isError are disjoint: a broken formula is something to
        report, not an empty cell to skip over."""
        value = CellValue(error_value)
        assert value.isError
        assert not value.isBlank

    def test_dash_is_not_an_error(self):
        assert not CellValue("-").isError

    def test_error_value_is_recognised_despite_whitespace(self):
        assert CellValue("  #REF!  ").isError

    def test_non_string_is_not_an_error(self):
        assert not CellValue(42).isError

    def test_whitespace_is_blank(self):
        assert CellValue("   ").isBlank

    def test_zero_is_not_blank(self):
        assert not CellValue(0).isBlank

    def test_false_is_not_blank(self):
        assert not CellValue(False).isBlank

    def test_text_is_not_blank(self):
        assert not CellValue("hello").isBlank

    def test_as_str_returns_value(self):
        assert CellValue("hello").as_str() == "hello"

    def test_as_str_stringifies_numbers(self):
        assert CellValue(42).as_str() == "42"

    def test_as_str_does_not_strip(self):
        assert CellValue("  hello  ").as_str() == "  hello  "

    def test_as_str_stripped_strips_whitespace(self):
        assert CellValue("  hello  ").as_str_stripped() == "hello"

    def test_as_str_stripped_stringifies_numbers(self):
        assert CellValue(42).as_str_stripped() == "42"

    def test_as_str_stripped_fallback_for_none(self):
        assert CellValue(None).as_str_stripped() == ""
        assert CellValue(None).as_str_stripped(fallback="FB") == "FB"

    def test_as_str_stripped_whitespace_only_is_empty_not_fallback(self):
        # fallback is for missing values only; a whitespace-only cell has a
        # value, it just strips to nothing (use isBlank to detect that case).
        assert CellValue("   ").as_str_stripped(fallback="FB") == ""

    def test_none_has_no_value(self):
        assert not CellValue(None).hasValue

    def test_placeholder_has_value(self):
        # hasValue is about cell emptiness, not placeholder semantics: a '-'
        # placeholder is a present value (use isBlank for the wider check).
        assert CellValue("-").hasValue

    def test_zero_has_value(self):
        assert CellValue(0).hasValue

    def test_as_str_fallback_for_none(self):
        assert CellValue(None).as_str() == ""
        assert CellValue(None).as_str(fallback="FB") == "FB"

    def test_as_date_from_date(self):
        d = date(2024, 12, 31)
        assert CellValue(d).as_date() == d

    def test_as_date_from_datetime(self):
        # Excel stores naive datetimes, so that is what a cell hands us.
        assert CellValue(datetime(2024, 6, 15, 10, 30)).as_date() == date(  # noqa: DTZ001
            2024, 6, 15
        )

    def test_as_date_from_iso_string(self):
        assert CellValue("2024-03-01").as_date() == date(2024, 3, 1)

    def test_as_date_from_none_raises(self):
        with pytest.raises(TypeError):
            CellValue(None).as_date()

    def test_from_cell_none_is_blank(self):
        assert CellValue.fromCell(None) == CellValue(None)

    def test_from_cell_takes_cell_value(self):
        class FakeCell:
            value = 42

        assert CellValue.fromCell(FakeCell()) == CellValue(42)  # type: ignore[arg-type]

    def test_from_cell_stringifies_rich_objects(self):
        class FakeCell:
            value: ClassVar[list[str]] = ["rich", "text"]

        assert CellValue.fromCell(FakeCell()).raw == "['rich', 'text']"  # type: ignore[arg-type]


class TestReaderValue:
    def test_known_name_yields_value(self, reader):
        v = reader.value("template_reporting_template_version")
        assert isinstance(v, CellValue)
        assert v.as_str() != ""

    def test_missing_name_yields_blank(self, reader):
        v = reader.value("this_does_not_exist_xyz")
        assert v.raw is None
        assert v.isBlank
        assert v.as_str(fallback="FB") == "FB"

    def test_period_start_as_date(self, reader):
        start_name = VSME_DEFAULTS["periods"][0]["start"]
        assert isinstance(reader.value(start_name).as_date(), date)


class TestUnusedDefinedNames:
    def test_unused_defined_names_populated_on_init(self, reader):
        assert len(reader.unused_defined_names) > 0

    def test_excluded_prefixes_absent(self, reader):
        from mireport.xlsx_template_reader._constants import (
            IGNORED_DEFINED_NAME_PREFIXES,
        )

        for dn in reader.unused_defined_names:
            assert not dn.name.startswith(IGNORED_DEFINED_NAME_PREFIXES)


@pytest.fixture(scope="module")
def bound(taxonomy):
    wb = loadExcelFromPathOrFileLike(SAMPLE)
    reader = WorkbookReader(wb, _results())
    try:
        yield WorkbookBinder(reader, taxonomy, _results()).bind()
    finally:
        wb.close()


class TestWorkbookBinder:
    def test_concept_map_populated(self, bound):
        assert len(bound.concept_map) > 0
        for dn, crm in bound.concept_map.items():
            assert isinstance(dn, DefinedName)
            assert isinstance(crm, XbrlConceptCellRangeMetadata)

    def test_tables_resolved(self, bound):
        assert len(bound.tables) > 0

    def test_table_ranges_consumed_from_concept_map(self, bound):
        for table_binding in bound.tables:
            for crm in table_binding.conceptRanges:
                assert crm.definedName not in bound.concept_map

    def test_table_units_consumed_from_unit_map(self, bound):
        for table_binding in bound.tables:
            for unit in table_binding.units:
                assert unit.concept not in bound.unit_map
