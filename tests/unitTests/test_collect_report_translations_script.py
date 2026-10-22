"""scripts/collect-report-translations.py on a small in-memory Translations sheet: what it keeps,
what it drops and says so, and what a rerun leaves alone."""

from __future__ import annotations

import importlib.util
from pathlib import Path
from types import ModuleType

import pytest
from babel.messages.catalog import Catalog
from babel.messages.pofile import read_po, write_po
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

SCRIPT = (
    Path(__file__).resolve().parents[2] / "scripts" / "collect-report-translations.py"
)

ENGLISH_COUNT = "This report contains {{ ns.fact_count }} XBRL facts ({{ facts | count }} unique facts)."
# label: (en, de, nl, ga)
ROWS = {
    "template_label_report_generated": (
        "Report generated",
        " Bericht erstellt\n\n",
        "Datum waarop rapport is gegenereerd",
        "Tuairisc ginte",
    ),
    "template_label_decimal_separator_document_information": (
        "Decimal separator",
        "Decimal separator",
        "Decimaal scheidingsteken",
        None,
    ),
    "template_label_entity_identifier_scheme_document_information": (
        "Entity Identifier Scheme",
        "System zur Unternehmenskennung",
        "Mij niet duidelijk waar deze wordt gebruikt",
        "Scéim Aitheantais Aonáin",
    ),
    "template_label_XBRL_facts_in_report": (
        ENGLISH_COUNT,
        "Dieser Bericht enthält {{ ns.fact_count }} XBRL-Fakten ({{facts|count}} eindeutige Fakten).",
        "Dit rapport bevat",
        "Tá fíricí XBRL {{ ns.fíric_comhaireamh }} sa tuarascáil seo.",
    ),
}


@pytest.fixture(scope="module")
def script() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "collect_report_translations_script", SCRIPT
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(autouse=True)
def noCorrections(script: ModuleType, monkeypatch: pytest.MonkeyPatch) -> None:
    """The real corrections are for the real template; each test sets its own."""
    monkeypatch.setattr(script, "CORRECTIONS", {})


@pytest.fixture(scope="module")
def workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Translations"
    ws.append(["Label ID", "Translation", "English", "German", "Dutch", "Irish"])
    ws.append([None, None, "en", "de", "nl", "ga"])
    ws.append([None, None, "English [en]", "German [de]", "Dutch [nl]", "Irish [ga]"])
    wb.defined_names["template_translations"] = DefinedName(
        "template_translations", attr_text="Translations!$C$3:$F$3"
    )
    for label, texts in ROWS.items():
        ws.append([label, texts[0], *texts])
        wb.defined_names[label] = DefinedName(
            label, attr_text=f"Translations!$B${ws.max_row}"
        )
    return wb


def catalog(directory: Path, language: str) -> Catalog:
    with (directory / f"{language}.po").open("rb") as f:
        return read_po(f)


def test_good_entries_are_kept_with_whitespace_stripped(
    script: ModuleType, workbook: Workbook, tmp_path: Path
) -> None:
    script.writeCatalogs(script.collect(workbook), tmp_path)
    de = catalog(tmp_path, "de")
    generated = de["Report generated"]
    assert generated.string == "Bericht erstellt"
    assert generated.auto_comments == ["template_label_report_generated"]
    assert catalog(tmp_path, "ga")["Report generated"].string == "Tuairisc ginte"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["de.po", "ga.po", "nl.po"]


def test_count_sentences_get_placeholders_and_a_fuzzy_one_count_form(
    script: ModuleType, workbook: Workbook, tmp_path: Path
) -> None:
    script.writeCatalogs(script.collect(workbook), tmp_path)
    de = catalog(tmp_path, "de")
    twoCounts = de[script.TWO_COUNTS]
    assert twoCounts.string == (
        "Dieser Bericht enthält %(factCount)s XBRL-Fakten "
        "(%(uniqueFactCount)s eindeutige Fakten)."
    )
    assert not twoCounts.fuzzy
    oneCount = de[script.ONE_COUNT]
    assert oneCount.string == "Dieser Bericht enthält %(factCount)s XBRL-Fakten."
    assert oneCount.fuzzy


def test_bad_entries_are_dropped_and_reported(
    script: ModuleType, workbook: Workbook, tmp_path: Path
) -> None:
    collected = script.collect(workbook)
    script.writeCatalogs(collected, tmp_path)
    dropped = [d for d in collected.dropped if not d.endswith("not in the template")]
    assert {d.split(":")[0]: d.split(": ")[1] for d in dropped} == {
        "de template_label_decimal_separator_document_information": "a copy of the English ('Decimal separator')",
        "ga template_label_decimal_separator_document_information": "blank ('')",
        "nl template_label_entity_identifier_scheme_document_information": "a translator's note, not a translation ('Mij niet duidelijk waar deze wordt gebruikt')",
        "nl template_label_XBRL_facts_in_report": "no fact count ('Dit rapport bevat')",
        "ga template_label_XBRL_facts_in_report": "an unknown placeholder ('Tá fíricí XBRL {{ ns.fíric_comhaireamh }} sa tuarascáil seo.')",
    }
    assert "Decimal separator" not in catalog(tmp_path, "de")
    nl = catalog(tmp_path, "nl")
    assert "Entity Identifier Scheme" not in nl
    assert script.TWO_COUNTS not in nl and script.ONE_COUNT not in nl
    assert collected.review == [
        "nl template_label_report_generated: long ('Datum waarop rapport is gegenereerd')"
    ]


def test_corrections_replace_the_template_and_are_reported(
    script: ModuleType,
    workbook: Workbook,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scheme = "template_label_entity_identifier_scheme_document_information"
    monkeypatch.setattr(
        script,
        "CORRECTIONS",
        {
            ("nl", scheme): "Identificatieschema entiteit",
            ("ga", "template_label_report_generated"): "Tuairisc ginte",
        },
    )
    collected = script.collect(workbook)
    script.writeCatalogs(collected, tmp_path)
    assert catalog(tmp_path, "nl")["Entity Identifier Scheme"].string == (
        "Identificatieschema entiteit"
    )
    assert collected.corrected == [
        f"nl {scheme}: 'Mij niet duidelijk waar deze wordt gebruikt' -> 'Identificatieschema entiteit'"
    ]
    assert (
        "ga template_label_report_generated: the template now has the correction"
        in (collected.review)
    )


def test_headers_are_filled_and_dated_only_on_change(
    script: ModuleType, workbook: Workbook, tmp_path: Path
) -> None:
    collected = script.collect(workbook)
    script.writeCatalogs(collected, tmp_path)
    de = catalog(tmp_path, "de")
    assert de.last_translator == "EFRAG Digital Template, Translations sheet"
    assert de.msgid_bugs_address.endswith("/issues")
    before = (tmp_path / "de.po").read_bytes()
    script.writeCatalogs(collected, tmp_path)
    assert (tmp_path / "de.po").read_bytes() == before


def test_a_rerun_keeps_hand_added_and_reviewed_entries(
    script: ModuleType, workbook: Workbook, tmp_path: Path
) -> None:
    script.writeCatalogs(script.collect(workbook), tmp_path)
    de = catalog(tmp_path, "de")
    de.add("Document Information", "Dokumentinformationen")
    de[script.ONE_COUNT].flags.discard("fuzzy")
    with (tmp_path / "de.po").open("wb") as f:
        write_po(f, de)

    script.writeCatalogs(script.collect(workbook), tmp_path)
    de = catalog(tmp_path, "de")
    assert de["Document Information"].string == "Dokumentinformationen"
    assert not de[script.ONE_COUNT].fuzzy
    assert de["Report generated"].string == "Bericht erstellt"
