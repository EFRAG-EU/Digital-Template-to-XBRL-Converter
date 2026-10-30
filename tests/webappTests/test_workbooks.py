"""Workbooks: each uploaded workbook parsed once, however often a request reads it."""

import gc
import weakref
from io import BytesIO
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook

from digital_converter_webapp.workbooks import Workbooks
from mireport.filesupport import FilelikeAndFileName


def _xlsx(title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_workbook_read_again_is_not_parsed_again(workbookParses):
    workbooks = Workbooks()
    upload = FilelikeAndFileName(_xlsx("A"), "report.xlsx")
    first = workbooks.load(upload)
    assert workbooks.load(upload) is first
    assert first.sheetnames == ["A"]
    assert workbookParses == [False]


def test_the_same_file_under_another_name_is_the_same_workbook(workbookParses):
    """This year's and last year's are told apart by what they hold, not their names."""
    workbooks = Workbooks()
    data = _xlsx("A")
    first = workbooks.load(FilelikeAndFileName(data, "this-year.xlsx"))
    assert workbooks.load(FilelikeAndFileName(data, "last-year.xlsx")) is first
    assert len(workbookParses) == 1


def test_different_files_with_one_name_are_different_workbooks(workbookParses):
    workbooks = Workbooks()
    first = workbooks.load(FilelikeAndFileName(_xlsx("A"), "report.xlsx"))
    second = workbooks.load(FilelikeAndFileName(_xlsx("B"), "report.xlsx"))
    assert (first.sheetnames, second.sheetnames) == (["A"], ["B"])
    assert len(workbookParses) == 2


def test_a_workbook_taken_is_the_one_loaded(workbookParses):
    workbooks = Workbooks()
    upload = FilelikeAndFileName(_xlsx("A"), "report.xlsx")
    loaded = workbooks.load(upload)
    assert workbooks.take(upload) is loaded
    assert len(workbookParses) == 1


def test_a_workbook_taken_is_let_go(workbookParses):
    """Its last reader takes it, so it is freed once read rather than at the end of the request."""
    workbooks = Workbooks()
    upload = FilelikeAndFileName(_xlsx("A"), "report.xlsx")
    workbooks.load(upload)
    taken = weakref.ref(workbooks.take(upload))
    gc.collect()
    assert taken() is None


def test_a_workbook_taken_without_being_loaded_is_parsed_and_not_kept(
    workbookParses,
):
    workbooks = Workbooks()
    upload = FilelikeAndFileName(_xlsx("A"), "report.xlsx")
    taken = weakref.ref(workbooks.take(upload))
    gc.collect()
    assert taken() is None
    assert len(workbookParses) == 1


def test_a_file_that_cant_be_parsed_is_tried_again(workbookParses):
    """So whoever reads it next says what is wrong with it in their own way."""
    workbooks = Workbooks()
    upload = FilelikeAndFileName(b"not an xlsx", "report.xlsx")
    for _ in range(2):
        with pytest.raises(BadZipFile):
            workbooks.load(upload)
    assert len(workbookParses) == 2
