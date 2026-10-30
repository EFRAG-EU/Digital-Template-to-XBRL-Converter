"""Results pages for the outcomes other than a clean conversion: a failed
conversion, an expired one, a report Arelle finds invalid, and one offered
migration. Arelle is faked wherever its outcome is the point, so these are
quick.
"""

import re
from io import BytesIO

import pytest
from cachelib.file import FileSystemCache
from openpyxl import Workbook
from playwright.sync_api import Page, expect
from tests.unitTests.arelle_results import cleanValidation

from digital_converter_webapp import create_app
from digital_converter_webapp import migration as migrationModule
from digital_converter_webapp.migration import MigrationOutcome
from mireport.arelle.report_info import ArelleReportProcessor

DOWNLOAD_BUTTONS = ("Download viewer", "Download XBRL Report", "Download xBRL-JSON")


@pytest.fixture
def noJsonValidation(monkeypatch):
    """The xBRL-JSON passes Arelle without Arelle running."""
    monkeypatch.setattr(
        ArelleReportProcessor,
        "validateXbrlJson",
        lambda self, document: cleanValidation(),
    )


def emptyWorkbook() -> dict:
    buffer = BytesIO()
    Workbook().save(buffer)
    return {
        "name": "empty.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "buffer": buffer.getvalue(),
    }


def test_failed_conversion_disables_downloads(page: Page, uploadAndConvert):
    uploadAndConvert(page, emptyWorkbook())

    expect(page.get_by_text("Technical conversion failed with ERRORS")).to_be_visible()
    expect(
        page.get_by_text("XBRL Report not validated due to earlier conversion failure")
    ).to_be_visible()
    expect(page.get_by_text("No files available for download")).to_be_visible()
    for name in DOWNLOAD_BUTTONS:
        expect(page.get_by_role("link", name=name)).to_have_attribute(
            "aria-disabled", "true"
        )


def test_unknown_conversion_is_expired(page: Page):
    response = page.goto("/conversions/no-such-conversion")

    assert response is not None and response.status == 404
    expect(page.get_by_text("Conversion expired")).to_be_visible()
    expect(page.get_by_text("No conversion results available.")).to_be_visible()


def test_validation_errors_disable_downloads(
    page: Page, uploadAndConvert, skipArelle, noJsonValidation
):
    """An Arelle ERROR turns the validation banner red and withholds the
    report: none of it can be downloaded."""
    uploadAndConvert(page)

    expect(
        page.get_by_text("The generated XBRL report has validation ERRORS")
    ).to_be_visible()
    expect(
        page.get_by_text("Arelle validation was not run (test stub)")
    ).to_be_visible()
    expect(page.get_by_text("No files available for download")).to_be_visible()
    for name in DOWNLOAD_BUTTONS:
        expect(page.get_by_role("link", name=name)).to_have_attribute(
            "aria-disabled", "true"
        )


@pytest.fixture(scope="module")
def migrationUrl(serve, tmp_path_factory) -> str:
    """An app with migration on. The flag is forced past ``create_app``, which
    only honours it with the migration tool installed: offering migration needs
    only the checks, faked here, not the tool."""
    app = create_app(
        {
            "TESTING": True,
            "SESSION_CACHELIB": FileSystemCache(
                str(tmp_path_factory.mktemp("migration-sessions"))
            ),
            "SECRET_KEY": "test-secret",
        }
    )
    app.config["ENABLE_MIGRATION"] = True
    return serve(app)


def test_optional_migration_is_offered(
    browser,
    browser_context_args,
    migrationUrl,
    uploadAndConvert,
    monkeypatch,
    skipArelle,
    noJsonValidation,
):
    monkeypatch.setattr(
        migrationModule,
        "doMigrationChecks",
        lambda conv, workbooks=None: (MigrationOutcome.MIGRATION_OPTIONAL, "1.2.0"),
    )
    context = browser.new_context(**{**browser_context_args, "base_url": migrationUrl})
    page = context.new_page()
    uploadAndConvert(page)

    offer = page.get_by_role("link", name="Migrate to Current Template Version")
    expect(offer).to_be_visible()
    expect(offer).to_have_attribute("href", re.compile(r"/migrationPage/"))
    context.close()
