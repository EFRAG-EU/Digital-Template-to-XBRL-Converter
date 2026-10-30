"""The steps run on a converted report, through the results page."""

import threading
import time
from io import BytesIO

import lxml.html
import pytest
from cachelib.file import FileSystemCache
from openpyxl import load_workbook
from tests.unitTests.arelle_results import cleanValidation

from digital_converter_webapp import create_app
from digital_converter_webapp.app_setup import ResultsLayout
from digital_converter_webapp.background import StepStore
from mireport.arelle.report_info import ArelleReportProcessor
from mireport.arelle.support import ArelleProcessingResult
from mireport.conversionresults import Step, StepState
from mireport.filesupport import FilelikeAndFileName

from .conftest import _test_config


@pytest.fixture
def gate():
    """The package's validation waits until this is open: it is, unless a test
    closes it. Opened again afterwards so no step is left waiting."""
    event = threading.Event()
    event.set()
    yield event
    event.set()


@pytest.fixture
def validations(monkeypatch, skipArelle, gate) -> list[str]:
    """Fake, counted validations: the package's carries an ERROR, the
    xBRL-JSON's is clean."""
    calls: list[str] = []
    fakePackage = ArelleReportProcessor.validateReportPackage

    def validateReportPackage(self, source, **kwargs):
        calls.append("package")
        assert gate.wait(timeout=60), "The test never let validation go ahead"
        return fakePackage(self, source, **kwargs)

    def validateXbrlJson(self, document):
        calls.append("xbrlJson")
        return cleanValidation()

    monkeypatch.setattr(
        ArelleReportProcessor, "validateReportPackage", validateReportPackage
    )
    monkeypatch.setattr(ArelleReportProcessor, "validateXbrlJson", validateXbrlJson)
    return calls


def test_results_page_shows_the_steps_outcome(client, upload_id, validations):
    conv_id = upload_id(client)

    page = client.get(f"/conversions/{conv_id}").get_data(as_text=True)

    assert validations == ["package", "xbrlJson"]
    assert "The generated XBRL report has validation ERRORS" in page
    assert "Arelle validation was not run (test stub)" in page


def brokenWorkbook(sample: bytes) -> bytes:
    """The sample workbook with no valid template version: an Excel parsing
    error."""
    workbook = load_workbook(BytesIO(sample))
    [(sheet, ref)] = workbook.defined_names[
        "template_reporting_template_version"
    ].destinations
    workbook[sheet][ref.replace("$", "")].value = "not a version"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# openpyxl drops the sample's extensions it can't read: no matter here.
@pytest.mark.filterwarnings("ignore:.* extension is not supported:UserWarning")
def test_a_workbook_with_errors_stops_the_conversion(
    app, client, upload, sample_xlsx, validations
):
    """No report is made from it, so none is validated or viewed: each step
    says it was not run, rather than waiting for ever."""
    response = upload(client, data=brokenWorkbook(sample_xlsx))
    conv_id = response.headers["Location"].rstrip("/").split("/")[-1]

    page = client.get(f"/conversions/{conv_id}").get_data(as_text=True)

    assert validations == []
    assert "Not finished yet" not in page
    assert "Not run: Digital Template (test.xlsx) has errors." in page
    assert page.count("Not run: the conversion failed.") == len(Step)
    assert "Inline XBRL report" in page
    assert "created (containing" not in page
    with app.app_context():
        assert StepStore.fromFlask().load(conv_id) == {}


def test_lost_steps_run_again(app, client, upload_id, validations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")
    with app.app_context():
        store = StepStore.fromFlask()
        store.forget(conv_id)

        page = client.get(f"/conversions/{conv_id}").get_data(as_text=True)

        assert validations == ["package", "xbrlJson"] * 2
        assert "The generated XBRL report has validation ERRORS" in page
        steps = store.load(conv_id)
        assert set(steps) == set(Step)
        assert [s.state for s in steps.values()] == [
            StepState.DONE,
            StepState.DONE,
            StepState.SKIPPED,
        ]


def test_finished_steps_are_not_run_again(client, upload_id, validations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")
    client.get(f"/conversions/{conv_id}")

    assert validations == ["package", "xbrlJson"]


def test_deleting_a_conversion_forgets_its_steps(app, client, upload_id, validations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    client.post(f"/delete/{conv_id}")

    with app.app_context():
        assert StepStore.fromFlask().load(conv_id) == {}


def test_status_of_a_finished_conversion(client, upload_id, validations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    status = client.get(f"/conversions/{conv_id}/status").get_json()

    assert status["complete"] is True
    assert "validation ERRORS" in status["html"]["validation-status"]
    assert 'id="downloads"' in status["html"]["downloads"]
    assert "Arelle validation was not run (test stub)" in status["html"]["messages"]
    messages = status["html"]["messages"]
    assert messages.lstrip().startswith('<div id="messages"')
    for section in ("workbook", "report", "package", "xbrlJson", "viewer"):
        assert f'id="section-{section}"' in messages


def test_flat_layout_lists_the_messages_in_one_table(tmp_path, upload_id, validations):
    client = create_app(_test_config(tmp_path, RESULTS_LAYOUT="flat")).test_client()
    conv_id = upload_id(client)

    page = client.get(f"/conversions/{conv_id}").get_data(as_text=True)

    [messages] = lxml.html.fromstring(page).xpath('//*[@id="messages"]')
    assert messages.tag == "table"
    assert not list(messages.iter("details"))
    assert "Excel Reference" in messages.xpath(".//th/text()")
    assert "Arelle validation was not run (test stub)" in messages.text_content()
    assert "Starting:" not in messages.text_content()

    dev = client.get(f"/conversions/{conv_id}?show_developer_messages=true")
    [messages] = lxml.html.fromstring(dev.data).xpath('//*[@id="messages"]')
    texts = [" ".join(t.split()) for t in messages.xpath(".//td[3]/text()")]
    package = texts.index("Starting: [XBRL report validation].")
    assert "Arelle validation was not run (test stub)" in texts[package:]
    assert any(
        t.startswith("Finished: [XBRL report validation] in ") for t in texts[package:]
    )


def test_status_of_an_unknown_conversion(client):
    response = client.get("/conversions/no-such-conversion/status")
    assert response.status_code == 404


@pytest.fixture(scope="module", params=list(ResultsLayout))
def backgroundClient(request, tmp_path_factory):
    """A client of an app that runs the steps in the background, as deployed, in
    each results layout."""
    app = create_app(
        {
            "TESTING": True,
            "SESSION_CACHELIB": FileSystemCache(
                str(tmp_path_factory.mktemp("background-sessions"))
            ),
            "SECRET_KEY": "test-secret",
            "RESULTS_LAYOUT": request.param,
        }
    )
    return app.test_client()


def waitUntilComplete(client, conv_id: str) -> dict:
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        status = client.get(f"/conversions/{conv_id}/status").get_json()
        if status["complete"]:
            return status
        time.sleep(0.1)
    raise AssertionError("The steps never finished")


def test_results_page_is_shown_before_validation_finishes(
    backgroundClient, upload_id, validations, gate
):
    gate.clear()
    conv_id = upload_id(backgroundClient)

    page = backgroundClient.get(f"/conversions/{conv_id}").get_data(as_text=True)

    assert 'data-state="pending"' in page
    assert f"/conversions/{conv_id}/status" in page
    assert "Validating the generated XBRL report" in page
    # The package's validation, in its section's header or under the banner.
    assert "waiting" in page or "in progress" in page
    status = backgroundClient.get(f"/conversions/{conv_id}/status").get_json()
    assert status["complete"] is False

    gate.set()
    status = waitUntilComplete(backgroundClient, conv_id)

    assert "validation ERRORS" in status["html"]["validation-status"]
    assert "Arelle validation was not run (test stub)" in status["html"]["messages"]
    page = backgroundClient.get(f"/conversions/{conv_id}").get_data(as_text=True)
    assert 'data-state="complete"' in page
    assert validations == ["package", "xbrlJson"]


REPORT_BUTTONS = {"Download viewer", "Download xBRL-JSON", "Download XBRL Report"}


def enabledButtons(html: str, elementId: str) -> set[str]:
    """The download buttons in the element that can be clicked."""
    [element] = lxml.html.fromstring(html).xpath(f'//*[@id="{elementId}"]')
    return {
        " ".join(link.text_content().split())
        for link in element.xpath('.//a[contains(@class, "download-handler")]')
        if link.get("aria-disabled") != "true"
    }


def test_downloads_wait_for_validation_on_both_pages(
    backgroundClient, upload_id, validations, gate
):
    """Only the upload can be had until validation has finished; then the
    report can't either, for it has errors."""
    gate.clear()
    conv_id = upload_id(backgroundClient)
    page = backgroundClient.get(f"/conversions/{conv_id}").get_data(as_text=True)
    listing = backgroundClient.get("/conversions/").get_data(as_text=True)

    assert enabledButtons(page, "downloads") == set()
    assert "once it has been validated" in page
    assert enabledButtons(listing, f"conversion-{conv_id}") == {"Download Excel"}
    [row] = lxml.html.fromstring(listing).xpath(f'//*[@id="conversion-{conv_id}"]')
    assert row.get("data-state") == "pending"
    rowStatus = backgroundClient.get(row.get("data-status-url")).get_json()
    assert rowStatus["complete"] is False
    assert list(rowStatus["html"]) == [f"conversion-{conv_id}"]

    gate.set()
    waitUntilComplete(backgroundClient, conv_id)

    rowStatus = backgroundClient.get(row.get("data-status-url")).get_json()
    [rowHtml] = rowStatus["html"].values()
    assert enabledButtons(rowHtml, f"conversion-{conv_id}") == {"Download Excel"}
    assert 'data-state="complete"' in rowHtml
    assert "data-status-url" not in rowHtml


def test_validated_report_can_be_downloaded_from_both_pages(
    client, upload_id, cleanValidations
):
    conv_id = upload_id(client)
    page = client.get(f"/conversions/{conv_id}").get_data(as_text=True)
    listing = client.get("/conversions/").get_data(as_text=True)

    assert enabledButtons(page, "downloads") == REPORT_BUTTONS
    assert enabledButtons(listing, f"conversion-{conv_id}") == REPORT_BUTTONS | {
        "Download Excel"
    }
    assert "data-status-url" not in listing


VIEWER = FilelikeAndFileName(fileContent=b"<html/>", filename="report_viewer.html")


@pytest.fixture
def cleanValidations(monkeypatch) -> list[str]:
    """Both validations pass and the viewer is built, without Arelle running.
    Returns the packages viewers were built from."""
    for name in ("validateReportPackage", "validateXbrlJson"):
        monkeypatch.setattr(
            ArelleReportProcessor,
            name,
            lambda *args, **kwargs: cleanValidation(),
        )
    built: list[str] = []

    def generateInlineViewer(self, source):
        built.append(source.filename)
        result = ArelleProcessingResult()
        result._viewer = VIEWER
        return result

    monkeypatch.setattr(
        ArelleReportProcessor, "generateInlineViewer", generateInlineViewer
    )
    return built


@pytest.mark.parametrize("ftype", ["zip", "json", "viewer"])
def test_validation_errors_withhold_the_report(client, upload_id, validations, ftype):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    for method in (client.head, client.get):
        response = method(f"/downloadFile/{conv_id}/{ftype}/")
        assert response.status_code == 409
    assert "Validation found errors" in response.get_json()["error"]
    status = client.get(f"/conversions/{conv_id}/status").get_json()
    assert status["downloadsAllowed"] is False
    assert "Validation found errors" in status["error"]


def test_validation_errors_withhold_the_served_viewer(client, upload_id, validations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    assert client.get(f"/viewer/{conv_id}/").status_code == 409


def test_validation_errors_leave_the_upload_downloadable(
    client, upload_id, validations
):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    assert client.get(f"/downloadFile/{conv_id}/excel/").status_code == 200


@pytest.mark.parametrize("ftype", ["zip", "json"])
def test_valid_report_is_downloadable(client, upload_id, cleanValidations, ftype):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    response = client.head(f"/downloadFile/{conv_id}/{ftype}/")

    assert response.status_code == 200
    assert response.headers["X-File-Ready"] == "true"
    status = client.get(f"/conversions/{conv_id}/status").get_json()
    assert status["downloadsAllowed"] is True
    assert status["error"] is None


def test_download_waits_for_validation(backgroundClient, upload_id, validations, gate):
    gate.clear()
    conv_id = upload_id(backgroundClient)
    backgroundClient.get(f"/conversions/{conv_id}")

    response = backgroundClient.head(f"/downloadFile/{conv_id}/json/")

    assert response.status_code == 202
    assert response.headers["X-File-Ready"] == "false"
    gate.set()
    waitUntilComplete(backgroundClient, conv_id)
    assert backgroundClient.head(f"/downloadFile/{conv_id}/json/").status_code == 409


def test_viewer_built_by_the_steps_is_served(client, upload_id, cleanValidations):
    conv_id = upload_id(client)
    client.get(f"/conversions/{conv_id}")

    response = client.get(f"/downloadFile/{conv_id}/viewer/")

    assert response.status_code == 200
    assert response.data == VIEWER.fileContent
    assert len(cleanValidations) == 1


def test_viewer_waits_for_its_step(
    backgroundClient, upload_id, cleanValidations, monkeypatch
):
    viewerGate = threading.Event()
    build = ArelleReportProcessor.generateInlineViewer

    def slowViewer(self, source):
        assert viewerGate.wait(timeout=60), "The test never let the viewer be built"
        return build(self, source)

    monkeypatch.setattr(ArelleReportProcessor, "generateInlineViewer", slowViewer)
    conv_id = upload_id(backgroundClient)
    backgroundClient.get(f"/conversions/{conv_id}")
    try:
        waitForStep(backgroundClient, conv_id, Step.VIEWER, StepState.RUNNING)

        assert (
            backgroundClient.head(f"/downloadFile/{conv_id}/json/").status_code == 200
        )
        assert (
            backgroundClient.head(f"/downloadFile/{conv_id}/viewer/").status_code == 202
        )
    finally:
        viewerGate.set()
    waitUntilComplete(backgroundClient, conv_id)
    assert backgroundClient.head(f"/downloadFile/{conv_id}/viewer/").status_code == 200
    assert len(cleanValidations) == 1


def waitForStep(client, conv_id: str, step: Step, state: StepState) -> None:
    deadline = time.monotonic() + 60
    with client.application.app_context():
        store = StepStore.fromFlask()
        while time.monotonic() < deadline:
            results = store.load(conv_id).get(step)
            if results is not None and results.state is state:
                return
            time.sleep(0.05)
    raise AssertionError(f"{step.label} never got to {state}")


def test_direct_download_waits_for_validation(
    backgroundClient, upload_id, validations, gate
):
    gate.clear()
    conv_id = upload_id(backgroundClient)
    backgroundClient.get(f"/conversions/{conv_id}")
    threading.Timer(0.5, gate.set).start()

    response = backgroundClient.get(f"/downloadFile/{conv_id}/json/")

    assert response.status_code == 409
    assert validations == ["package", "xbrlJson"]


def test_direct_download_stops_waiting(
    backgroundClient, upload_id, validations, gate, monkeypatch
):
    monkeypatch.setitem(
        backgroundClient.application.config, "DOWNLOAD_WAIT_SECONDS", 0.2
    )
    gate.clear()
    conv_id = upload_id(backgroundClient)
    backgroundClient.get(f"/conversions/{conv_id}")

    response = backgroundClient.get(f"/downloadFile/{conv_id}/json/")

    assert response.status_code == 202
    assert response.headers["X-File-Ready"] == "false"
    assert "still being validated" in response.get_json()["status"]
    gate.set()
    waitUntilComplete(backgroundClient, conv_id)
