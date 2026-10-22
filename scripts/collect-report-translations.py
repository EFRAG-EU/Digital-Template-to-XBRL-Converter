"""Collect the translations of the report's fixed strings from a Digital Template.

Templates carry reviewed translations of the strings the converter writes into every report
(the Document Information keys, the headings, the fact-count sentence) on their Translations
sheet. This writes them into one gettext catalog per language, merged into the catalogs already
there so that hand-added entries survive. The checked-in catalogs are the source of truth:
review their git diff after a run, and rerun when EFRAG ships new translations.
"""

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from babel.messages.catalog import Catalog
from babel.messages.pofile import read_po, write_po
from openpyxl import Workbook, load_workbook

CATALOGS = (
    Path(__file__).resolve().parents[1]
    / "src/mireport/data/disclosures/translations/vsme"
)

ONE_COUNT = "This report contains %(factCount)s XBRL facts."
TWO_COUNTS = (
    "This report contains %(factCount)s XBRL facts (%(uniqueFactCount)s unique facts)."
)

# Each label on the Translations sheet that the report uses, and the msgid it translates.
LABEL_MSGIDS = {
    "template_label_entity_name_document_information": "Entity Name",
    "template_label_entity_identifier_document_information": "Entity Identifier",
    "template_label_entity_identifier_scheme_document_information": "Entity Identifier Scheme",
    "template_label_report_currency_document_information": "Report currency",
    "template_label_decimal_separator_document_information": "Decimal separator",
    "template_label_report_period_document_information": "Report period",
    "template_label_report_generated": "Report generated",
    "template_label_disclosures_related_to_reporting_period": "All disclosures are related to the reporting period above unless specified otherwise.",
    "template_label_table_of_contents": "Table of contents",
    "template_label_reporting_period": "Reporting period:",
    "template_label_XBRL_facts_in_report": TWO_COUNTS,
    "template_label_introduction": "Introduction",
}

# (language, label, reason): entries that pass every check but are not translations.
KNOWN_BAD = [
    (
        "nl",
        "template_label_entity_identifier_scheme_document_information",
        "a translator's note, not a translation",
    ),
]

# (language, label): our reviewed fix for a template's translation, used in its place. Drop
# one once the template has it; the summary says so.
CORRECTIONS = {
    ("da", "template_label_introduction"): "Indledning",
    (
        "de",
        "template_label_XBRL_facts_in_report",
    ): "Dieser Bericht enthält %(factCount)s XBRL-Fakten (%(uniqueFactCount)s eindeutige Fakten).",
    ("fr", "template_label_introduction"): "Introduction",
    ("it", "template_label_reporting_period"): "Periodo di riferimento:",
    (
        "nl",
        "template_label_entity_identifier_document_information",
    ): "Identificatiecode entiteit",
    ("nl", "template_label_report_currency_document_information"): "Rapportagevaluta",
    ("pl", "template_label_reporting_period"): "Okres sprawozdawczy:",
    (
        "sl",
        "template_label_entity_identifier_document_information",
    ): "Identifikacijska številka podjetja",
}

HEADERS = {
    "project": "VSME Digital Template",
    "msgid_bugs_address": "https://github.com/EFRAG-EU/Digital-Template-to-XBRL-Converter/issues",
    "last_translator": "EFRAG Digital Template, Translations sheet",
    "language_team": "EFRAG Digital Template, Translations sheet",
}

PLACEHOLDERS = {
    re.compile(r"\{\{\s*ns\.fact_count\s*\}\}"): "%(factCount)s",
    re.compile(r"\{\{\s*facts\s*\|\s*count\s*\}\}"): "%(uniqueFactCount)s",
}
ENGLISH = "en"
# A translation this many times longer than its English is listed for review, not dropped.
LONG = 2


@dataclass(frozen=True)
class Entry:
    label: str
    msgid: str
    string: str
    fuzzy: bool = False


@dataclass
class Collected:
    entries: dict[str, list[Entry]] = field(default_factory=dict)
    dropped: list[str] = field(default_factory=list)
    corrected: list[str] = field(default_factory=list)
    review: list[str] = field(default_factory=list)
    templateVersion: str | None = None


def languageColumns(workbook: Workbook) -> dict[str, int]:
    """Each language's column: the codes in the row above the template_translations range."""
    ((sheet, cells),) = workbook.defined_names["template_translations"].destinations
    ws = workbook[sheet]
    ((first, *_),) = ws[cells.replace("$", "")]
    return {
        code.strip(): cell.column
        for cell in ws[first.row - 1]
        if cell.column >= first.column and isinstance(code := cell.value, str)
    }


def labelRows(workbook: Workbook) -> dict[str, int]:
    """The Translations sheet row of each label in LABEL_MSGIDS that the template has."""
    rows = {}
    for label in LABEL_MSGIDS:
        if (name := workbook.defined_names.get(label)) is None:
            continue
        for sheet, cells in name.destinations:
            rows[label] = workbook[sheet][cells.replace("$", "")].row
    return rows


def templateVersion(workbook: Workbook) -> str | None:
    """The template's own version, recorded in each catalog's header."""
    if (
        name := workbook.defined_names.get("template_reporting_template_version")
    ) is None:
        return None
    for sheet, cells in name.destinations:
        return cleaned(workbook[sheet][cells.replace("$", "")].value) or None
    return None


def cleaned(text: object) -> str:
    text = str(text or "").strip()
    for placeholder, substitute in PLACEHOLDERS.items():
        text = placeholder.sub(substitute, text)
    return text


def problem(language: str, label: str, text: str, english: str) -> str | None:
    """Why this translation can't be used, or None if it can."""
    msgid = LABEL_MSGIDS[label]
    if not text:
        return "blank"
    if text.casefold() in {english.casefold(), msgid.casefold()}:
        return "a copy of the English"
    if "{{" in text or "}}" in text:
        return "an unknown placeholder"
    if msgid == TWO_COUNTS and "%(factCount)s" not in text:
        return "no fact count"
    return next(
        (why for lang, lab, why in KNOWN_BAD if (lang, lab) == (language, label)), None
    )


def entriesFor(label: str, text: str) -> list[Entry]:
    """A count sentence fills the form its placeholders match. The one-count form is also
    derived from the two-count one, as a fuzzy entry, since EFRAG has no one-count form."""
    msgid = LABEL_MSGIDS[label]
    if msgid != TWO_COUNTS:
        return [Entry(label, msgid, text)]
    if "%(uniqueFactCount)s" not in text:
        return [Entry(label, ONE_COUNT, text, fuzzy=True)]
    entries = [Entry(label, TWO_COUNTS, text)]
    if " (" in text:
        oneCount = text.partition(" (")[0] + "."
        entries.append(Entry(label, ONE_COUNT, oneCount, fuzzy=True))
    return entries


def collect(workbook: Workbook) -> Collected:
    columns = languageColumns(workbook)
    rows = labelRows(workbook)
    ws = workbook["Translations"]
    collected = Collected(templateVersion=templateVersion(workbook))
    for label in sorted(LABEL_MSGIDS.keys() - rows.keys()):
        collected.dropped.append(f"{label}: not in the template")
    for language, column in columns.items():
        if language == ENGLISH:
            continue
        entries = collected.entries.setdefault(language, [])
        for label, row in rows.items():
            english = cleaned(ws.cell(row, columns[ENGLISH]).value)
            text = cleaned(ws.cell(row, column).value)
            if (correction := CORRECTIONS.get((language, label))) is not None:
                if correction == text and not problem(language, label, text, english):
                    collected.review.append(
                        f"{language} {label}: the template now has the correction"
                    )
                elif correction != text:
                    collected.corrected.append(
                        f"{language} {label}: {text!r} -> {correction!r}"
                    )
                entries.extend(entriesFor(label, correction))
                continue
            if why := problem(language, label, text, english):
                collected.dropped.append(f"{language} {label}: {why} ({text!r})")
                continue
            if len(text) > LONG * len(english):
                collected.review.append(f"{language} {label}: long ({text!r})")
            entries.extend(entriesFor(label, text))
    return collected


def mergeInto(
    path: Path, language: str, entries: list[Entry], version: str | None = None
) -> Catalog:
    """The catalog at ``path`` (or a new one) with these entries set. Entries it has that
    are not among them are kept; one whose translation is unchanged keeps its flags, so a
    reviewed fuzzy entry stays reviewed. The revision date moves only when something
    changes."""
    changed = not path.exists()
    if path.exists():
        with path.open("rb") as f:
            catalog = read_po(f, locale=language)
    else:
        catalog = Catalog(
            locale=language,
            header_comment="# Collected by scripts/collect-report-translations.py.",
            fuzzy=False,
        )
    for name, value in {**HEADERS, "version": version or catalog.version}.items():
        if getattr(catalog, name) != value:
            setattr(catalog, name, value)
            changed = True
    for entry in entries:
        flags = {"fuzzy"} if entry.fuzzy else set()
        if (existing := catalog.get(entry.msgid)) is None:
            catalog.add(entry.msgid, entry.string, flags=flags)
            existing = catalog[entry.msgid]
            changed = True
        elif existing.string != entry.string:
            existing.string = entry.string
            existing.flags = flags
            changed = True
        existing.auto_comments = [entry.label]
    if changed:
        catalog.revision_date = datetime.now(UTC)
    return catalog


def writeCatalogs(collected: Collected, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    for language, entries in collected.entries.items():
        path = directory / f"{language}.po"
        catalog = mergeInto(path, language, entries, collected.templateVersion)
        with path.open("wb") as f:
            write_po(f, catalog, width=0)


def summary(collected: Collected) -> list[str]:
    lines = []
    for language, entries in sorted(collected.entries.items()):
        fuzzy = sum(e.fuzzy for e in entries)
        dropped = sum(d.startswith(f"{language} ") for d in collected.dropped)
        lines.append(
            f"{language}: {len(entries) - fuzzy} kept, {fuzzy} fuzzy, {dropped} dropped"
        )
    lines += [f"Dropped {d}" for d in collected.dropped]
    lines += [f"Corrected {c}" for c in collected.corrected]
    lines += [f"For review {r}" for r in collected.review]
    return lines


def createArgParser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Collect the report's translated fixed strings from a Digital "
        "Template's Translations sheet into gettext catalogs, one per language."
    )
    parser.add_argument("template", type=Path, help="A current Digital Template")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=CATALOGS,
        help=f"Where the <language>.po catalogs are (default: {CATALOGS})",
    )
    return parser


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")  # the summary quotes every language
    args = createArgParser().parse_args()
    collected = collect(load_workbook(args.template, data_only=True))
    writeCatalogs(collected, args.output_dir)
    for line in summary(collected):
        print(line)
    print("Now compile them: scripts/compile-report-translations.py")


if __name__ == "__main__":
    main()
